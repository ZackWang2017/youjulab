"""从互联网拉取 SH / SZ / HK / US 证券代码并写入 securities 表。

数据源：
  - SH / SZ：新浪财经 Market_Center API（A 股 + B 股，以及 ETF / LOF 场内基金，记为 security_type=基金）
  - HK：港交所官方 ListOfSecurities.xlsx
  - US：NASDAQ Trader nasdaqtraded.txt（NASDAQ / NYSE / NYSE American / Arca / BATS / IEX 的股票与 ETF，
    ETF 记为 security_type=基金，其余为 股票）

用法（在 mcp-server 目录下）：
  ../.venv/bin/python financial-data-service/repository/securities_repository.py
  ../.venv/bin/python financial-data-service/repository/securities_repository.py --markets SH,SZ
  ../.venv/bin/python financial-data-service/repository/securities_repository.py --markets US
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Iterable

import openpyxl

# 允许从 mcp-server 目录直接运行本脚本
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from db.connection import get_connection

USER_AGENT = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"
SINA_BASE = "https://vip.stock.finance.sina.com.cn/quotes_service/api/json_v2.php"
SINA_PAGE_SIZE = 80
# A 股 + B 股（900xxx.SH 以美元计价，200xxx.SZ 以港币计价）
SINA_NODES = {"SH": ("sh_a", "sh_b"), "SZ": ("sz_a", "sz_b")}
# 沪深场内基金（ETF / LOF）：节点不分市场，按 symbol 前缀 sh/sz 归属
SINA_FUND_NODES = ("etf_hq_fund", "lof_hq_fund")
HKEX_LIST_URL = (
    "https://www.hkex.com.hk/eng/services/trading/securities/"
    "securitieslists/ListOfSecurities.xlsx"
)
NASDAQ_TRADED_URL = "https://www.nasdaqtrader.com/dynamic/SymDir/nasdaqtraded.txt"
# Listing Exchange 列 → 展示用交易所简称（写进名称里方便辨认，与 holdings 参考实现一致）
US_EXCHANGE_LABEL = {"A": "NYSE MKT", "N": "NYSE", "P": "NYSE ARCA", "Q": "NASDAQ", "Z": "BATS", "V": "IEX"}
_US_SYMBOL_RE = re.compile(r"^[A-Z]{1,6}$")  # 跳过权证 / 优先股 / 单位等带 . $ 的代码
DEFAULT_MARKETS = ("SH", "SZ", "HK", "US")
SUPPORTED_MARKETS = frozenset(DEFAULT_MARKETS)
SECURITY_TYPE_STOCK = "股票"
SECURITY_TYPE_FUND = "基金"
NAME_MAX_LEN = 128  # securities.name 为 varchar(128)


@dataclass(frozen=True, slots=True)
class SecurityRecord:
    ticker_code: str
    name: str
    security_type: str = SECURITY_TYPE_STOCK


def _http_get(url: str, *, referer: str = "https://finance.sina.com.cn/", timeout: int = 30) -> bytes:
    req = urllib.request.Request(
        url,
        headers={"User-Agent": USER_AGENT, "Referer": referer},
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read()


def _http_get_with_retry(url: str, *, retries: int = 3, pause: float = 1.0, **kwargs) -> bytes:
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            return _http_get(url, **kwargs)
        except (urllib.error.URLError, TimeoutError) as exc:
            last_error = exc
            if attempt < retries:
                time.sleep(pause * attempt)
    raise RuntimeError(f"request failed after {retries} retries: {url}") from last_error


def format_ticker(code: str, market: str) -> str:
    code = str(code).strip()
    if market == "HK":
        code = code.zfill(5)
    return f"{code}.{market}"


def fetch_sina_market(market: str) -> list[SecurityRecord]:
    records: list[SecurityRecord] = []
    for node in SINA_NODES[market]:
        records.extend(_fetch_sina_node(node, market=market))
    records.extend(r for r in _fetch_sina_funds() if r.ticker_code.endswith(f".{market}"))
    return records


_sina_funds_cache: list[SecurityRecord] | None = None


def _fetch_sina_funds() -> list[SecurityRecord]:
    """沪深 ETF / LOF 清单，一次运行内只拉一遍（SH / SZ 共用）。

    新浪给的基金简称有三分之一被截成「纳指」「食品」这类 2~4 字，再用腾讯行情的全称覆盖。
    """
    global _sina_funds_cache
    if _sina_funds_cache is None:
        funds: list[SecurityRecord] = []
        for node in SINA_FUND_NODES:
            funds.extend(_fetch_sina_node(node, security_type=SECURITY_TYPE_FUND))
        _sina_funds_cache = _enrich_names_from_tencent(funds)
    return _sina_funds_cache


TENCENT_QUOTE_URL = "https://qt.gtimg.cn/q="
_TENCENT_BATCH = 60


def _enrich_names_from_tencent(records: list[SecurityRecord]) -> list[SecurityRecord]:
    """用腾讯实时行情的名称字段（第 2 列）替换更短的新浪简称。请求失败时保留原名。"""
    symbol_of = {f"{r.ticker_code.rsplit('.', 1)[1].lower()}{r.ticker_code.split('.', 1)[0]}": r for r in records}
    symbols = list(symbol_of)
    full_names: dict[str, str] = {}
    for i in range(0, len(symbols), _TENCENT_BATCH):
        chunk = symbols[i:i + _TENCENT_BATCH]
        try:
            body = _http_get_with_retry(TENCENT_QUOTE_URL + ",".join(chunk), referer="https://gu.qq.com/")
        except RuntimeError:
            continue
        for line in body.decode("gbk", errors="replace").splitlines():
            line = line.strip().rstrip(";")
            if not line.startswith("v_") or "=" not in line:
                continue
            symbol, _, payload = line[2:].partition("=")
            fields = payload.strip().strip('"').split("~")
            if len(fields) > 1 and fields[1].strip():
                full_names[symbol] = fields[1].strip()
    out: list[SecurityRecord] = []
    for symbol, record in symbol_of.items():
        full = full_names.get(symbol, "")
        if len(full) > len(record.name):
            record = SecurityRecord(record.ticker_code, full[:NAME_MAX_LEN], record.security_type)
        out.append(record)
    return out


def _market_from_symbol(symbol: str) -> str | None:
    prefix = symbol[:2].lower()
    return {"sh": "SH", "sz": "SZ", "bj": "BJ"}.get(prefix)


def _fetch_sina_node(node: str, *, market: str | None = None, security_type: str = SECURITY_TYPE_STOCK) -> list[SecurityRecord]:
    """拉取一个新浪行情节点。market 为空时按每条的 symbol 前缀（sh/sz）判断市场。"""
    count_url = f"{SINA_BASE}/Market_Center.getHQNodeStockCount?node={node}"
    total = int(json.loads(_http_get_with_retry(count_url)))
    pages = math.ceil(total / SINA_PAGE_SIZE)

    records: list[SecurityRecord] = []
    for page in range(1, pages + 1):
        data_url = (
            f"{SINA_BASE}/Market_Center.getHQNodeData"
            f"?page={page}&num={SINA_PAGE_SIZE}&sort=symbol&asc=1&node={node}"
        )
        batch = json.loads(_http_get_with_retry(data_url))
        for item in batch:
            code = str(item.get("code", "")).strip()
            name = str(item.get("name", "")).strip()
            if not code or not name or name == "FAILED":
                continue
            item_market = market or _market_from_symbol(str(item.get("symbol", "")))
            if not item_market:
                continue
            records.append(
                SecurityRecord(
                    ticker_code=format_ticker(code, item_market),
                    name=name[:NAME_MAX_LEN],
                    security_type=security_type,
                )
            )
    return records


def fetch_hk_equities() -> list[SecurityRecord]:
    content = _http_get_with_retry(
        HKEX_LIST_URL,
        referer="https://www.hkex.com.hk/",
        timeout=60,
    )
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    records: list[SecurityRecord] = []
    for row in worksheet.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        name = str(row[1]).strip() if row[1] else ""
        category = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        if not code or not name or category != "Equity":
            continue
        records.append(
            SecurityRecord(
                ticker_code=format_ticker(code, "HK"),
                name=name,
            )
        )
    return records


def parse_nasdaq_traded(text: str) -> list[SecurityRecord]:
    """解析 nasdaqtraded.txt（竖线分隔）。ETF 列为 Y 记为基金，其余记为股票。"""
    lines = text.splitlines()
    if not lines:
        return []
    header = [h.strip() for h in lines[0].split("|")]
    try:
        idx_symbol = header.index("Symbol")
        idx_name = header.index("Security Name")
        idx_exch = header.index("Listing Exchange")
        idx_etf = header.index("ETF")
        idx_test = header.index("Test Issue")
    except ValueError as exc:
        raise RuntimeError(f"nasdaqtraded.txt unexpected header: {header}") from exc
    needed = max(idx_symbol, idx_name, idx_exch, idx_etf, idx_test)

    records: list[SecurityRecord] = []
    for raw in lines[1:]:
        if not raw.strip() or raw.startswith("File Creation Time"):
            continue
        parts = raw.split("|")
        if len(parts) <= needed:
            continue
        symbol = parts[idx_symbol].strip().upper()
        name = " ".join(parts[idx_name].split())
        if not symbol or not name or not _US_SYMBOL_RE.match(symbol):
            continue
        if parts[idx_test].strip().upper() == "Y":
            continue
        label = US_EXCHANGE_LABEL.get(parts[idx_exch].strip().upper())
        suffix = f" [{label}]" if label and label not in name else ""
        name = name[: NAME_MAX_LEN - len(suffix)].rstrip() + suffix
        is_etf = parts[idx_etf].strip().upper() == "Y"
        records.append(
            SecurityRecord(
                ticker_code=format_ticker(symbol, "US"),
                name=name,
                security_type=SECURITY_TYPE_FUND if is_etf else SECURITY_TYPE_STOCK,
            )
        )
    return records


def fetch_us_securities() -> list[SecurityRecord]:
    content = _http_get_with_retry(NASDAQ_TRADED_URL, referer="https://www.nasdaqtrader.com/", timeout=60)
    return parse_nasdaq_traded(content.decode("utf-8", errors="replace"))


def fetch_market(market: str) -> list[SecurityRecord]:
    market = market.upper()
    if market in SINA_NODES:
        return fetch_sina_market(market)
    if market == "HK":
        return fetch_hk_equities()
    if market == "US":
        return fetch_us_securities()
    raise ValueError(f"unsupported market: {market}")


def save_securities(records: Iterable[SecurityRecord]) -> int:
    rows = [
        (item.ticker_code, item.name, item.security_type)
        for item in records
    ]
    if not rows:
        return 0

    sql = """
        INSERT INTO securities (ticker_code, name, security_type)
        VALUES (%s, %s, %s::security_type_enum)
        ON CONFLICT (ticker_code) DO UPDATE SET
            name = EXCLUDED.name,
            security_type = EXCLUDED.security_type,
            updated_at = NOW()
    """
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.executemany(sql, rows)
    return len(rows)


def sync_securities(*, markets: Iterable[str] = DEFAULT_MARKETS) -> dict:
    summary: dict[str, int | str] = {"saved_total": 0}
    for market in markets:
        market = market.upper()
        records = fetch_market(market)
        saved = save_securities(records)
        summary[market] = saved
        summary["saved_total"] = int(summary["saved_total"]) + saved
    return summary


def _parse_markets(raw: str) -> tuple[str, ...]:
    markets = tuple(part.strip().upper() for part in raw.split(",") if part.strip())
    invalid = [m for m in markets if m not in SUPPORTED_MARKETS]
    if invalid:
        raise ValueError(f"unsupported markets: {', '.join(invalid)}")
    return markets


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Sync SH/SZ/HK/US securities into PostgreSQL")
    parser.add_argument(
        "--markets",
        default=",".join(DEFAULT_MARKETS),
        help="comma-separated markets: SH,SZ,HK,US",
    )
    args = parser.parse_args(argv)

    try:
        markets = _parse_markets(args.markets)
        summary = sync_securities(markets=markets)
    except Exception as exc:  # noqa: BLE001
        print(f"sync failed: {exc}", file=sys.stderr)
        return 1

    for market in markets:
        print(f"{market}: {summary[market]} records")
    print(f"total saved/updated: {summary['saved_total']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
